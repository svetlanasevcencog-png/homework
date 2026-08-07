#!/usr/bin/env python3
"""Generate a PowerPoint deck on QA harness techniques — CGI brand + diagrams."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

CGI_RED = RGBColor(0xE3, 0x19, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF7, 0xF7, 0xF7)
CARD = RGBColor(0xF0, 0xF0, 0xF0)
CHARCOAL = RGBColor(0x1A, 0x1A, 0x1A)
DARK = RGBColor(0x2D, 0x2D, 0x2D)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)
LIGHT_LINE = RGBColor(0xD9, 0xD9, 0xD9)
SOFT_RED = RGBColor(0xFB, 0xE8, 0xEB)
TOTAL = 10

W, H = Inches(13.333), Inches(7.5)


def set_run(run, size=18, bold=False, color=CHARCOAL, font="Arial"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_top_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CGI_RED
    bar.line.fill.background()


def add_text_box(slide, left, top, width, height, text, size=18, bold=False,
                 color=CHARCOAL, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_footer(slide, page, total=TOTAL):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.015)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_LINE
    line.line.fill.background()
    add_text_box(slide, Inches(0.6), Inches(7.15), Inches(3), Inches(0.3),
                 "CGI", size=11, bold=True, color=CGI_RED)
    add_text_box(slide, Inches(3.5), Inches(7.15), Inches(7), Inches(0.3),
                 "QA Harness — Techniques Overview", size=10, color=MUTED)
    add_text_box(slide, Inches(11.5), Inches(7.15), Inches(1.2), Inches(0.3),
                 f"{page} / {total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def add_bullets(slide, left, top, width, height, items, size=14, color=DARK, spacing=6):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run(run, size=size, color=color)
    return box


def add_card(slide, left, top, width, height, title, body_lines, title_color=CGI_RED):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.fill.background()
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    strip.fill.solid()
    strip.fill.fore_color.rgb = CGI_RED
    strip.line.fill.background()
    add_text_box(slide, left + Inches(0.25), top + Inches(0.2), width - Inches(0.4),
                 Inches(0.35), title, size=13, bold=True, color=title_color)
    add_bullets(slide, left + Inches(0.25), top + Inches(0.6), width - Inches(0.4),
                height - Inches(0.75), body_lines, size=12, color=DARK, spacing=4)


def blank_slide(prs, page, bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    add_top_bar(slide)
    add_footer(slide, page)
    return slide


def section_header(slide, title, subtitle=None):
    add_text_box(slide, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5),
                 title, size=26, bold=True, color=CHARCOAL)
    if subtitle:
        add_text_box(slide, Inches(0.6), Inches(0.82), Inches(12), Inches(0.3),
                     subtitle, size=13, color=MUTED)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.18), Inches(1.4), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = CGI_RED
    line.line.fill.background()


def node(slide, left, top, width, height, text, fill=CARD, text_color=CHARCOAL, size=11):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = CGI_RED if fill == CARD or fill == SOFT_RED else fill
    shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    try:
        shape.text_frame.auto_size = None
    except Exception:
        pass
    run = tf.paragraphs[0].add_run()
    run.text = text
    set_run(run, size=size, bold=True, color=text_color)
    tf.paragraphs[0].space_before = Pt(6)
    return shape


def arrow_r(slide, left, top, width=Inches(0.28), height=Inches(0.22)):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    a.fill.solid()
    a.fill.fore_color.rgb = CGI_RED
    a.line.fill.background()
    return a


def arrow_d(slide, left, top, width=Inches(0.22), height=Inches(0.28)):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    a.fill.solid()
    a.fill.fore_color.rgb = CGI_RED
    a.line.fill.background()
    return a


def diamond(slide, left, top, width, height, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = SOFT_RED
    shape.line.color.rgb = CGI_RED
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = text
    set_run(run, size=11, bold=True, color=CGI_RED)
    tf.paragraphs[0].space_before = Pt(22)
    return shape


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 Title
    s = blank_slide(prs, 1)
    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), H)
    panel.fill.solid()
    panel.fill.fore_color.rgb = CGI_RED
    panel.line.fill.background()
    add_text_box(s, Inches(0.9), Inches(1.8), Inches(4), Inches(0.4),
                 "CGI", size=22, bold=True, color=CGI_RED)
    add_text_box(s, Inches(0.9), Inches(2.4), Inches(11), Inches(0.35),
                 "TECHNIQUES OVERVIEW", size=13, bold=True, color=MUTED)
    add_text_box(s, Inches(0.9), Inches(2.9), Inches(11.5), Inches(1.0),
                 "Agent-Driven QA Harness", size=40, bold=True, color=CHARCOAL)
    add_text_box(s, Inches(0.9), Inches(4.1), Inches(10), Inches(0.9),
                 "Cursor agents · Playwright conventions · CI guardrails\n"
                 "How the harness turns tickets into reliable automated tests",
                 size=16, color=MUTED)
    add_text_box(s, Inches(0.9), Inches(5.5), Inches(10), Inches(0.35),
                 "Insights you can act on", size=14, bold=True, color=CGI_RED)
    add_text_box(s, Inches(0.9), Inches(6.2), Inches(10), Inches(0.3),
                 "Focus: engineering techniques — not the product under test",
                 size=12, color=MUTED)

    # 2 Agenda
    s = blank_slide(prs, 2)
    section_header(s, "Agenda", "What we will cover")
    items = [
        ("01", "Problem & approach"),
        ("02", "Architecture"),
        ("03", "Agents & entry points"),
        ("04", "Test stack & skills"),
        ("05", "Conventions & guardrails"),
        ("06", "Happy path"),
        ("07", "Heal on red & CI"),
        ("08", "Key takeaways"),
    ]
    for i, (num, label) in enumerate(items):
        col, row = i % 2, i // 2
        left = Inches(0.7) + col * Inches(6.1)
        top = Inches(1.55) + row * Inches(1.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.8), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        num_box = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.2), top + Inches(0.22),
            Inches(0.75), Inches(0.55))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = CGI_RED
        num_box.line.fill.background()
        tf = num_box.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tf.paragraphs[0].add_run()
        run.text = num
        set_run(run, size=14, bold=True, color=WHITE)
        tf.paragraphs[0].space_before = Pt(8)
        add_text_box(s, left + Inches(1.15), top + Inches(0.28), Inches(4.4), Inches(0.5),
                     label, size=16, bold=True, color=CHARCOAL)

    # 3 Problem
    s = blank_slide(prs, 3)
    section_header(s, "Problem & approach", "Why a harness — not just a test suite")
    add_card(s, Inches(0.6), Inches(1.5), Inches(5.9), Inches(5.0), "THE CHALLENGE", [
        "Ticket → tests is slow and inconsistent",
        "UI drift breaks locators silently",
        "Agents can invent paths or weaken asserts",
        "Flakes and leftover data erode trust",
        "Humans must stay in the loop for risk",
    ])
    add_card(s, Inches(6.8), Inches(1.5), Inches(5.9), Inches(5.0), "THE APPROACH", [
        "Coordinator delegates to specialist agents",
        "Skills encode repeatable procedures",
        "Constitution: MUST / SHOULD / WON'T",
        "Hooks block weakened assertions",
        "Human gates on plan review & merge",
        "Reliability eval before declaring Done",
    ])

    # 4 Architecture
    s = blank_slide(prs, 4)
    section_header(s, "Architecture", "From trigger to external systems")
    layers = [
        ("Triggers", "Jira tickets · gaps · red CI · cron"),
        ("Coordinator", "QA Orchestrator (+ headless CI runner)"),
        ("Agents", "test-writer · triage · bug-reporter"),
        ("Skills", "Analyzer, explore, heal, cleanup, eval…"),
        ("Test stack", "Features → specs → pages → fixtures"),
        ("Guardrails", "Constitution · hooks · conventions · humans"),
        ("CI / external", "GitHub Actions · Jira · project under test · MCPs"),
    ]
    for i, (title, desc) in enumerate(layers):
        top = Inches(1.45) + i * Inches(0.72)
        pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), top, Inches(2.5), Inches(0.55))
        pill.fill.solid()
        pill.fill.fore_color.rgb = CGI_RED if i % 2 == 0 else CHARCOAL
        pill.line.fill.background()
        tf = pill.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tf.paragraphs[0].add_run()
        run.text = title
        set_run(run, size=12, bold=True, color=WHITE)
        tf.paragraphs[0].space_before = Pt(10)
        add_text_box(s, Inches(3.5), top + Inches(0.08), Inches(9), Inches(0.45),
                     desc, size=14, color=DARK)

    # 5 Agents + entry points
    s = blank_slide(prs, 5)
    section_header(s, "Agents & entry points", "Narrow roles — coordinator never does everything alone")
    entries = [
        ("Jira ticket", "DS-* AC → Gherkin plan"),
        ("Coverage gap", "Live UI vs specs"),
        ("Red CI run", "triage → heal or bug"),
        ("Daily cron", "Headless backlog"),
    ]
    for i, (title, desc) in enumerate(entries):
        left = Inches(0.55) + i * Inches(3.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.45), Inches(3.05), Inches(1.55))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        add_text_box(s, left + Inches(0.15), Inches(1.65), Inches(2.75), Inches(0.4),
                     f"{i + 1}. {title}", size=13, bold=True, color=CHARCOAL, align=PP_ALIGN.CENTER)
        add_text_box(s, left + Inches(0.15), Inches(2.15), Inches(2.75), Inches(0.55),
                     desc, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    agents = [
        ("test-writer", ["Plan → specs + POMs", "pom-conventions & api-cleanup", "Opens feat/* PR"]),
        ("triage", ["Read-only on red runs", "Classifies app bug vs drift", "Never patches code"]),
        ("bug-reporter", ["Files / links Jira bugs", "Only after real bug", "Never self-heals"]),
    ]
    for i, (name, lines) in enumerate(agents):
        left = Inches(0.55) + i * Inches(4.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(3.25), Inches(4.0), Inches(3.15))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        header = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(3.25), Inches(4.0), Inches(0.55))
        header.fill.solid()
        header.fill.fore_color.rgb = CGI_RED
        header.line.fill.background()
        add_text_box(s, left + Inches(0.2), Inches(3.35), Inches(3.6), Inches(0.4),
                     name, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_bullets(s, left + Inches(0.35), Inches(4.05), Inches(3.4), Inches(2.0),
                    lines, size=13, color=DARK, spacing=10)

    # 6 Stack + skills
    s = blank_slide(prs, 6)
    section_header(s, "Test stack & skills", "Strict layering + reusable procedures")
    layers = [
        ("features/*.feature.md", "Gherkin plans — human-reviewed", True),
        ("tests/ds*.spec.ts", "Intent + expect() — no inline locators", False),
        ("pages/*.page.ts", "getByRole + actions — no assertions", False),
        ("helpers · fixtures · auth", "Shared flows · trackProgram · storageState", False),
    ]
    for i, (name, desc, accent) in enumerate(layers):
        top = Inches(1.4) + i * Inches(0.7)
        block = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), top, Inches(12), Inches(0.6))
        block.fill.solid()
        block.fill.fore_color.rgb = CGI_RED if accent else CARD
        block.line.fill.background()
        nc = WHITE if accent else CHARCOAL
        dc = RGBColor(0xFF, 0xE0, 0xE5) if accent else MUTED
        add_text_box(s, Inches(1.0), top + Inches(0.1), Inches(4.5), Inches(0.4),
                     name, size=12, bold=True, color=nc)
        add_text_box(s, Inches(5.6), top + Inches(0.1), Inches(6.8), Inches(0.4),
                     desc, size=12, color=dc)
    skills = [
        ("jira-ticket-analyzer", "AC → Gherkin plans"),
        ("explore-and-generate", "Coverage gaps"),
        ("ci-failure-triage", "Logs → diagnosis"),
        ("self-heal", "POM drift only"),
        ("jira-bug-reporter", "Structured bugs"),
        ("suite-reliability-eval", "eval-report.md"),
    ]
    for i, (name, desc) in enumerate(skills):
        col, row = i % 3, i // 3
        left = Inches(0.55) + col * Inches(4.2)
        top = Inches(4.4) + row * Inches(1.15)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(4.0), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), Inches(1.05))
        strip.fill.solid()
        strip.fill.fore_color.rgb = CGI_RED
        strip.line.fill.background()
        add_text_box(s, left + Inches(0.25), top + Inches(0.2), Inches(3.5), Inches(0.3),
                     name, size=12, bold=True, color=CHARCOAL)
        add_text_box(s, left + Inches(0.25), top + Inches(0.55), Inches(3.5), Inches(0.3),
                     desc, size=11, color=MUTED)

    # 7 Conventions + guardrails
    s = blank_slide(prs, 7)
    section_header(s, "Conventions & guardrails", "Machines enforce; humans decide risk")
    cols = [
        ("MUST", CGI_RED, [
            "Role-based locators",
            "Web-first assertions",
            "storageState via setup",
            "Self-cleaning data",
            "One tag per test()",
        ]),
        ("SHOULD", CHARCOAL, [
            "Faker factories",
            "Enums for UI strings",
            "JSDoc on POM actions",
        ]),
        ("WON'T", RGBColor(0x8B, 0x15, 0x28), [
            "CSS / XPath as default",
            "waitForTimeout",
            "Weaken assertions",
            "Heal a real app bug",
        ]),
    ]
    for i, (title, color, items) in enumerate(cols):
        left = Inches(0.5) + i * Inches(4.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.4), Inches(4.0), Inches(2.85))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        badge = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.25), Inches(1.6), Inches(1.6), Inches(0.4))
        badge.fill.solid()
        badge.fill.fore_color.rgb = color
        badge.line.fill.background()
        add_text_box(s, left + Inches(0.25), Inches(1.65), Inches(1.6), Inches(0.35),
                     title, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_bullets(s, left + Inches(0.3), Inches(2.2), Inches(3.5), Inches(1.9),
                    items, size=12, color=DARK, spacing=6)
    add_card(s, Inches(0.5), Inches(4.45), Inches(6.0), Inches(2.1), "ENFORCED", [
        "Constitution always on · hook blocks weakened expect()",
        "≤5 delegations/task · ≤5 heals/run",
        "Generation gate + main test check",
    ])
    add_card(s, Inches(6.8), Inches(4.45), Inches(6.0), Inches(2.1), "HUMAN GATES", [
        "Review Gherkin plan before coding",
        "Confidence <5 ask · 5–7 approve",
        "Merge never automated · never invent",
    ])

    # 8 Happy path
    s = blank_slide(prs, 8)
    section_header(s, "Happy path: ticket → green", "One ticket = one PR; humans approve the plan")
    steps = [
        ("1", "Trigger"), ("2", "Analyze"), ("3", "Plan"), ("4", "Review"),
        ("5", "Write"), ("6", "Run"), ("7", "PR"), ("8", "Eval"),
    ]
    for i, (num, title) in enumerate(steps):
        left = Inches(0.45) + i * Inches(1.6)
        node(s, left, Inches(1.7), Inches(1.35), Inches(1.1), f"{num}\n{title}",
             fill=CGI_RED if i in (0, 3, 7) else CARD,
             text_color=WHITE if i in (0, 3, 7) else CHARCOAL, size=12)
        if i < 7:
            arrow_r(s, left + Inches(1.35), Inches(2.1), Inches(0.22), Inches(0.2))
    add_card(s, Inches(0.6), Inches(3.3), Inches(5.9), Inches(3.1), "HUMAN IN THE LOOP", [
        "Plan review before any specs are written",
        "PR merge never automated by agents",
        "Confidence gate on uncertain edits",
    ])
    add_card(s, Inches(6.8), Inches(3.3), Inches(5.9), Inches(3.1), "QUALITY GATES", [
        "CI job test green after PR open",
        "eval-report.md refreshed before Done",
        "Empty/red checks → stop",
    ])

    # 9 Heal + CI
    s = blank_slide(prs, 9)
    section_header(s, "Heal on red & CI", "Classify first — right signal at the right gate")
    outcomes = [
        ("Real app bug", CGI_RED, ["bug-reporter → Jira", "Never self-heal"]),
        ("Locator drift", CHARCOAL, ["self-heal → patch pages/", "Assertions unchanged · ≤5"]),
        ("Inconclusive", RGBColor(0x8B, 0x15, 0x28), ["Stop & escalate", "No silent guess"]),
    ]
    for i, (title, color, items) in enumerate(outcomes):
        left = Inches(0.55) + i * Inches(4.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.45), Inches(4.0), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        topbar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.45), Inches(4.0), Inches(0.1))
        topbar.fill.solid()
        topbar.fill.fore_color.rgb = color
        topbar.line.fill.background()
        add_text_box(s, left + Inches(0.25), Inches(1.75), Inches(3.5), Inches(0.35),
                     title, size=14, bold=True, color=color)
        add_bullets(s, left + Inches(0.25), Inches(2.25), Inches(3.5), Inches(1.2),
                    items, size=12, color=DARK, spacing=8)
    rows = [
        ("Pull request → main", "@smoke", "Fast confidence on every PR"),
        ("Push → main", "@sanity", "Broader post-merge check"),
        ("Manual dispatch", "@regression", "Full suite on demand"),
        ("Daily cron", "QA orchestrator", "Backlog via cursor-agent"),
    ]
    headers = ["Trigger", "Suite / job", "Intent"]
    lefts = [Inches(0.7), Inches(4.5), Inches(8.0)]
    widths = [Inches(3.6), Inches(3.3), Inches(4.5)]
    for h, left, w in zip(headers, lefts, widths):
        add_text_box(s, left, Inches(3.9), w, Inches(0.3), h, size=11, bold=True, color=CGI_RED)
    for i, (trig, suite, intent) in enumerate(rows):
        top = Inches(4.3) + i * Inches(0.55)
        bg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), top, Inches(12.1), Inches(0.48))
        bg.fill.solid()
        bg.fill.fore_color.rgb = CARD if i % 2 == 0 else OFF_WHITE
        bg.line.fill.background()
        for text, left, w in zip([trig, suite, intent], lefts, widths):
            add_text_box(s, left, top + Inches(0.08), w, Inches(0.35), text, size=12, color=CHARCOAL)

    # 10 Takeaways
    s = blank_slide(prs, 10)
    section_header(s, "Key takeaways", "Techniques you can reuse")
    takeaways = [
        "Separate coordination from specialist work — orchestrator delegates, agents stay narrow.",
        "Encode process as skills; encode safety as constitution + hooks, not hope.",
        "Layer tests: plans → assertions → locators → data.",
        "Classify before healing — drift → POM; product bugs → Jira.",
        "Humans own plan approval and merge; agents own speed inside those rails.",
    ]
    for i, t in enumerate(takeaways):
        top = Inches(1.5) + i * Inches(0.7)
        num = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), top, Inches(0.45), Inches(0.45))
        num.fill.solid()
        num.fill.fore_color.rgb = CGI_RED
        num.line.fill.background()
        tf = num.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tf.paragraphs[0].add_run()
        run.text = str(i + 1)
        set_run(run, size=12, bold=True, color=WHITE)
        tf.paragraphs[0].space_before = Pt(6)
        add_text_box(s, Inches(1.4), top + Inches(0.05), Inches(11), Inches(0.45),
                     t, size=13, color=DARK)
    add_text_box(s, Inches(0.9), Inches(5.3), Inches(11), Inches(0.4),
                 "Questions?", size=22, bold=True, color=CGI_RED)
    add_text_box(s, Inches(0.9), Inches(5.85), Inches(11), Inches(0.4),
                 "docs/harness-diagram.md · .cursor/rules/ · .cursor/skills/",
                 size=13, color=MUTED)

    out = "/Users/svetlana.sevcenco/Formation/homework/docs/QA-Harness-Techniques.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
