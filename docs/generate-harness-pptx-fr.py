#!/usr/bin/env python3
"""Generate a French PowerPoint deck on QA harness techniques — CGI brand + diagrams."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

CGI_RED = RGBColor(0xE3, 0x19, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF7, 0xF7, 0xF7)
CARD = RGBColor(0xF0, 0xF0, 0xF0)
CHARCOAL = RGBColor(0x1A, 0x1A, 0x1A)
DARK = RGBColor(0x2D, 0x2D, 0x2D)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)
LIGHT_LINE = RGBColor(0xD9, 0xD9, 0xD9)
SOFT_RED = RGBColor(0xFB, 0xE8, 0xEB)
TOTAL = 10

W, H = Inches(13.333), Inches(7.5)


def set_run(run, size=18, bold=False, color=CHARCOAL, font="Arial"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_top_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CGI_RED
    bar.line.fill.background()


def add_text_box(slide, left, top, width, height, text, size=18, bold=False,
                 color=CHARCOAL, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_footer(slide, page, total=TOTAL):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(7.05), Inches(12.1), Inches(0.015)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_LINE
    line.line.fill.background()
    add_text_box(slide, Inches(0.6), Inches(7.15), Inches(3), Inches(0.3),
                 "CGI", size=11, bold=True, color=CGI_RED)
    add_text_box(slide, Inches(3.5), Inches(7.15), Inches(7), Inches(0.3),
                 "Harness QA — Aperçu des techniques", size=10, color=MUTED)
    add_text_box(slide, Inches(11.5), Inches(7.15), Inches(1.2), Inches(0.3),
                 f"{page} / {total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def add_bullets(slide, left, top, width, height, items, size=14, color=DARK, spacing=6):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run(run, size=size, color=color)
    return box


def add_card(slide, left, top, width, height, title, body_lines, title_color=CGI_RED):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.fill.background()
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    strip.fill.solid()
    strip.fill.fore_color.rgb = CGI_RED
    strip.line.fill.background()
    add_text_box(slide, left + Inches(0.25), top + Inches(0.2), width - Inches(0.4),
                 Inches(0.35), title, size=13, bold=True, color=title_color)
    add_bullets(slide, left + Inches(0.25), top + Inches(0.6), width - Inches(0.4),
                height - Inches(0.75), body_lines, size=12, color=DARK, spacing=4)


def blank_slide(prs, page, bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    add_top_bar(slide)
    add_footer(slide, page)
    return slide


def section_header(slide, title, subtitle=None):
    add_text_box(slide, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5),
                 title, size=26, bold=True, color=CHARCOAL)
    if subtitle:
        add_text_box(slide, Inches(0.6), Inches(0.82), Inches(12), Inches(0.3),
                     subtitle, size=13, color=MUTED)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.18), Inches(1.4), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = CGI_RED
    line.line.fill.background()


def node(slide, left, top, width, height, text, fill=CARD, text_color=CHARCOAL, size=11):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = CGI_RED if fill == CARD or fill == SOFT_RED else fill
    shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    try:
        shape.text_frame.auto_size = None
    except Exception:
        pass
    run = tf.paragraphs[0].add_run()
    run.text = text
    set_run(run, size=size, bold=True, color=text_color)
    tf.paragraphs[0].space_before = Pt(6)
    return shape


def arrow_r(slide, left, top, width=Inches(0.28), height=Inches(0.22)):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    a.fill.solid()
    a.fill.fore_color.rgb = CGI_RED
    a.line.fill.background()
    return a


def arrow_d(slide, left, top, width=Inches(0.22), height=Inches(0.28)):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
    a.fill.solid()
    a.fill.fore_color.rgb = CGI_RED
    a.line.fill.background()
    return a


def diamond(slide, left, top, width, height, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = SOFT_RED
    shape.line.color.rgb = CGI_RED
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = text
    set_run(run, size=11, bold=True, color=CGI_RED)
    tf.paragraphs[0].space_before = Pt(22)
    return shape


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 Title
    s = blank_slide(prs, 1)
    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), H)
    panel.fill.solid()
    panel.fill.fore_color.rgb = CGI_RED
    panel.line.fill.background()
    add_text_box(s, Inches(0.9), Inches(1.8), Inches(4), Inches(0.4),
                 "CGI", size=22, bold=True, color=CGI_RED)
    add_text_box(s, Inches(0.9), Inches(2.4), Inches(11), Inches(0.35),
                 "APERÇU DES TECHNIQUES", size=13, bold=True, color=MUTED)
    add_text_box(s, Inches(0.9), Inches(2.9), Inches(11.5), Inches(1.0),
                 "Harness QA piloté par agents", size=40, bold=True, color=CHARCOAL)
    add_text_box(s, Inches(0.9), Inches(4.1), Inches(10), Inches(0.9),
                 "Agents Cursor · conventions Playwright · garde-fous CI\n"
                 "Comment le harness transforme les tickets en tests automatisés fiables",
                 size=16, color=MUTED)
    add_text_box(s, Inches(0.9), Inches(5.5), Inches(10), Inches(0.35),
                 "Des insights actionnables", size=14, bold=True, color=CGI_RED)
    add_text_box(s, Inches(0.9), Inches(6.2), Inches(10), Inches(0.3),
                 "Focus : techniques d'ingénierie — pas le produit sous test",
                 size=12, color=MUTED)

    # 2 Agenda
    s = blank_slide(prs, 2)
    section_header(s, "Agenda", "Ce que nous allons couvrir")
    items = [
        ("01", "Problème et approche"),
        ("02", "Architecture"),
        ("03", "Agents et points d'entrée"),
        ("04", "Stack de tests et skills"),
        ("05", "Conventions et garde-fous"),
        ("06", "Happy path"),
        ("07", "Heal on red et CI"),
        ("08", "Points clés"),
    ]
    for i, (num, label) in enumerate(items):
        col, row = i % 2, i // 2
        left = Inches(0.7) + col * Inches(6.1)
        top = Inches(1.55) + row * Inches(1.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.8), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        num_box = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.2), top + Inches(0.22),
            Inches(0.75), Inches(0.55))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = CGI_RED
        num_box.line.fill.background()
        tf = num_box.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tf.paragraphs[0].add_run()
        run.text = num
        set_run(run, size=14, bold=True, color=WHITE)
        tf.paragraphs[0].space_before = Pt(8)
        add_text_box(s, left + Inches(1.15), top + Inches(0.28), Inches(4.4), Inches(0.5),
                     label, size=16, bold=True, color=CHARCOAL)

    # 3 Problem
    s = blank_slide(prs, 3)
    section_header(s, "Problème et approche", "Pourquoi un harness — pas seulement une suite de tests")
    add_card(s, Inches(0.6), Inches(1.5), Inches(5.9), Inches(5.0), "LE DÉFI", [
        "Ticket → tests : lent et incohérent",
        "La dérive UI casse les locators en silence",
        "Les agents peuvent inventer des chemins ou affaiblir les asserts",
        "Flakes et données résiduelles érodent la confiance",
        "Les humains doivent rester dans la boucle pour le risque",
    ])
    add_card(s, Inches(6.8), Inches(1.5), Inches(5.9), Inches(5.0), "L'APPROCHE", [
        "Le coordinateur délègue aux agents spécialistes",
        "Les skills encodent des procédures répétables",
        "Constitution : MUST / SHOULD / WON'T",
        "Les hooks bloquent les assertions affaiblies",
        "Portes humaines sur revue de plan et merge",
        "Éval. de fiabilité avant de déclarer Done",
    ])

    # 4 Architecture
    s = blank_slide(prs, 4)
    section_header(s, "Architecture", "Du déclencheur aux systèmes externes")
    layers = [
        ("Déclencheurs", "Tickets Jira · gaps · CI rouge · cron"),
        ("Coordinateur", "QA Orchestrator (+ runner CI headless)"),
        ("Agents", "test-writer · triage · bug-reporter"),
        ("Skills", "Analyzer, explore, heal, cleanup, eval…"),
        ("Stack de tests", "Features → specs → pages → fixtures"),
        ("Garde-fous", "Constitution · hooks · conventions · humains"),
        ("CI / externe", "GitHub Actions · Jira · projet sous test · MCPs"),
    ]
    for i, (title, desc) in enumerate(layers):
        top = Inches(1.45) + i * Inches(0.72)
        pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), top, Inches(2.5), Inches(0.55))
        pill.fill.solid()
        pill.fill.fore_color.rgb = CGI_RED if i % 2 == 0 else CHARCOAL
        pill.line.fill.background()
        tf = pill.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tf.paragraphs[0].add_run()
        run.text = title
        set_run(run, size=12, bold=True, color=WHITE)
        tf.paragraphs[0].space_before = Pt(10)
        add_text_box(s, Inches(3.5), top + Inches(0.08), Inches(9), Inches(0.45),
                     desc, size=14, color=DARK)

    # 5 Agents + entry points
    s = blank_slide(prs, 5)
    section_header(s, "Agents et points d'entrée", "Rôles étroits — le coordinateur ne fait jamais tout seul")
    entries = [
        ("Ticket Jira", "CA DS-* → plan Gherkin"),
        ("Gap couverture", "UI live vs specs"),
        ("Run CI rouge", "triage → heal ou bug"),
        ("Cron quotidien", "Backlog headless"),
    ]
    for i, (title, desc) in enumerate(entries):
        left = Inches(0.55) + i * Inches(3.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.45), Inches(3.05), Inches(1.55))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        add_text_box(s, left + Inches(0.15), Inches(1.65), Inches(2.75), Inches(0.4),
                     f"{i + 1}. {title}", size=13, bold=True, color=CHARCOAL, align=PP_ALIGN.CENTER)
        add_text_box(s, left + Inches(0.15), Inches(2.15), Inches(2.75), Inches(0.55),
                     desc, size=11, color=MUTED, align=PP_ALIGN.CENTER)
    agents = [
        ("test-writer", ["Plan → specs + POMs", "pom-conventions & api-cleanup", "Ouvre PR feat/*"]),
        ("triage", ["Lecture seule sur runs rouges", "Classe bug app vs dérive", "Ne patch jamais le code"]),
        ("bug-reporter", ["Crée / lie bugs Jira", "Seulement si vrai bug", "Jamais de self-heal"]),
    ]
    for i, (name, lines) in enumerate(agents):
        left = Inches(0.55) + i * Inches(4.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(3.25), Inches(4.0), Inches(3.15))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        header = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(3.25), Inches(4.0), Inches(0.55))
        header.fill.solid()
        header.fill.fore_color.rgb = CGI_RED
        header.line.fill.background()
        add_text_box(s, left + Inches(0.2), Inches(3.35), Inches(3.6), Inches(0.4),
                     name, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_bullets(s, left + Inches(0.35), Inches(4.05), Inches(3.4), Inches(2.0),
                    lines, size=13, color=DARK, spacing=10)

    # 6 Stack + skills
    s = blank_slide(prs, 6)
    section_header(s, "Stack de tests et skills", "Couches strictes + procédures réutilisables")
    layers = [
        ("features/*.feature.md", "Plans Gherkin — revue humaine", True),
        ("tests/ds*.spec.ts", "Intention + expect() — pas de locators inline", False),
        ("pages/*.page.ts", "getByRole + actions — pas d'assertions", False),
        ("helpers · fixtures · auth", "Flux partagés · trackProgram · storageState", False),
    ]
    for i, (name, desc, accent) in enumerate(layers):
        top = Inches(1.4) + i * Inches(0.7)
        block = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), top, Inches(12), Inches(0.6))
        block.fill.solid()
        block.fill.fore_color.rgb = CGI_RED if accent else CARD
        block.line.fill.background()
        nc = WHITE if accent else CHARCOAL
        dc = RGBColor(0xFF, 0xE0, 0xE5) if accent else MUTED
        add_text_box(s, Inches(1.0), top + Inches(0.1), Inches(4.5), Inches(0.4),
                     name, size=12, bold=True, color=nc)
        add_text_box(s, Inches(5.6), top + Inches(0.1), Inches(6.8), Inches(0.4),
                     desc, size=12, color=dc)
    skills = [
        ("jira-ticket-analyzer", "CA → plans Gherkin"),
        ("explore-and-generate", "Gaps de couverture"),
        ("ci-failure-triage", "Logs → diagnostic"),
        ("self-heal", "Dérive POM uniquement"),
        ("jira-bug-reporter", "Bugs structurés"),
        ("suite-reliability-eval", "eval-report.md"),
    ]
    for i, (name, desc) in enumerate(skills):
        col, row = i % 3, i // 3
        left = Inches(0.55) + col * Inches(4.2)
        top = Inches(4.4) + row * Inches(1.15)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(4.0), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), Inches(1.05))
        strip.fill.solid()
        strip.fill.fore_color.rgb = CGI_RED
        strip.line.fill.background()
        add_text_box(s, left + Inches(0.25), top + Inches(0.2), Inches(3.5), Inches(0.3),
                     name, size=12, bold=True, color=CHARCOAL)
        add_text_box(s, left + Inches(0.25), top + Inches(0.55), Inches(3.5), Inches(0.3),
                     desc, size=11, color=MUTED)

    # 7 Conventions + guardrails
    s = blank_slide(prs, 7)
    section_header(s, "Conventions et garde-fous", "Les machines appliquent ; les humains décident du risque")
    cols = [
        ("MUST", CGI_RED, [
            "Locators par rôle",
            "Assertions web-first",
            "storageState via setup",
            "Données auto-nettoyées",
            "Un tag par test()",
        ]),
        ("SHOULD", CHARCOAL, [
            "Factories Faker",
            "Enums pour UI",
            "JSDoc sur POM",
        ]),
        ("WON'T", RGBColor(0x8B, 0x15, 0x28), [
            "CSS / XPath par défaut",
            "waitForTimeout",
            "Affaiblir les asserts",
            "Healer un vrai bug",
        ]),
    ]
    for i, (title, color, items) in enumerate(cols):
        left = Inches(0.5) + i * Inches(4.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.4), Inches(4.0), Inches(2.85))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        badge = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.25), Inches(1.6), Inches(1.6), Inches(0.4))
        badge.fill.solid()
        badge.fill.fore_color.rgb = color
        badge.line.fill.background()
        add_text_box(s, left + Inches(0.25), Inches(1.65), Inches(1.6), Inches(0.35),
                     title, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_bullets(s, left + Inches(0.3), Inches(2.2), Inches(3.5), Inches(1.9),
                    items, size=12, color=DARK, spacing=6)
    add_card(s, Inches(0.5), Inches(4.45), Inches(6.0), Inches(2.1), "APPLIQUÉ", [
        "Constitution toujours active · hook bloque expect() affaiblis",
        "≤5 délégations/tâche · ≤5 heals/run",
        "Porte de génération + check test sur main",
    ])
    add_card(s, Inches(6.8), Inches(4.45), Inches(6.0), Inches(2.1), "PORTES HUMAINES", [
        "Revue du plan Gherkin avant code",
        "Confiance <5 demander · 5–7 approbation",
        "Merge jamais automatisé · ne jamais inventer",
    ])

    # 8 Happy path
    s = blank_slide(prs, 8)
    section_header(s, "Happy path : ticket → vert", "Un ticket = une PR ; les humains approuvent le plan")
    steps = [
        ("1", "Déclencher"), ("2", "Analyser"), ("3", "Plan"), ("4", "Revue"),
        ("5", "Écrire"), ("6", "Exécuter"), ("7", "PR"), ("8", "Éval."),
    ]
    for i, (num, title) in enumerate(steps):
        left = Inches(0.45) + i * Inches(1.6)
        node(s, left, Inches(1.7), Inches(1.35), Inches(1.1), f"{num}\n{title}",
             fill=CGI_RED if i in (0, 3, 7) else CARD,
             text_color=WHITE if i in (0, 3, 7) else CHARCOAL, size=12)
        if i < 7:
            arrow_r(s, left + Inches(1.35), Inches(2.1), Inches(0.22), Inches(0.2))
    add_card(s, Inches(0.6), Inches(3.3), Inches(5.9), Inches(3.1), "HUMAIN DANS LA BOUCLE", [
        "Revue du plan avant toute écriture de specs",
        "Merge de PR jamais automatisé par les agents",
        "Porte de confiance sur les edits incertains",
    ])
    add_card(s, Inches(6.8), Inches(3.3), Inches(5.9), Inches(3.1), "PORTES QUALITÉ", [
        "Job CI test vert après ouverture de PR",
        "eval-report.md rafraîchi avant Done",
        "Checks vides/rouges → stop",
    ])

    # 9 Heal + CI
    s = blank_slide(prs, 9)
    section_header(s, "Heal on red et CI", "Classifier d'abord — le bon signal à la bonne porte")
    outcomes = [
        ("Vrai bug app", CGI_RED, ["bug-reporter → Jira", "Jamais de self-heal"]),
        ("Dérive de locator", CHARCOAL, ["self-heal → patch pages/", "Assertions inchangées · ≤5"]),
        ("Inconclusif", RGBColor(0x8B, 0x15, 0x28), ["Stop & escalader", "Pas de guess silencieux"]),
    ]
    for i, (title, color, items) in enumerate(outcomes):
        left = Inches(0.55) + i * Inches(4.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.45), Inches(4.0), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD
        card.line.fill.background()
        topbar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.45), Inches(4.0), Inches(0.1))
        topbar.fill.solid()
        topbar.fill.fore_color.rgb = color
        topbar.line.fill.background()
        add_text_box(s, left + Inches(0.25), Inches(1.75), Inches(3.5), Inches(0.35),
                     title, size=14, bold=True, color=color)
        add_bullets(s, left + Inches(0.25), Inches(2.25), Inches(3.5), Inches(1.2),
                    items, size=12, color=DARK, spacing=8)
    rows = [
        ("Pull request → main", "@smoke", "Confiance rapide sur chaque PR"),
        ("Push → main", "@sanity", "Contrôle post-merge plus large"),
        ("Dispatch manuel", "@regression", "Suite complète à la demande"),
        ("Cron quotidien", "QA orchestrator", "Backlog via cursor-agent"),
    ]
    headers = ["Déclencheur", "Suite / job", "Intention"]
    lefts = [Inches(0.7), Inches(4.5), Inches(8.0)]
    widths = [Inches(3.6), Inches(3.3), Inches(4.5)]
    for h, left, w in zip(headers, lefts, widths):
        add_text_box(s, left, Inches(3.9), w, Inches(0.3), h, size=11, bold=True, color=CGI_RED)
    for i, (trig, suite, intent) in enumerate(rows):
        top = Inches(4.3) + i * Inches(0.55)
        bg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), top, Inches(12.1), Inches(0.48))
        bg.fill.solid()
        bg.fill.fore_color.rgb = CARD if i % 2 == 0 else OFF_WHITE
        bg.line.fill.background()
        for text, left, w in zip([trig, suite, intent], lefts, widths):
            add_text_box(s, left, top + Inches(0.08), w, Inches(0.35), text, size=12, color=CHARCOAL)

    # 10 Takeaways
    s = blank_slide(prs, 10)
    section_header(s, "Points clés", "Techniques réutilisables")
    takeaways = [
        "Séparer coordination et travail spécialisé — l'orchestrator délègue, les agents restent étroits.",
        "Encoder le processus en skills ; la sécurité en constitution + hooks, pas en espoir.",
        "Empiler les tests : plans → assertions → locators → données.",
        "Classifier avant de healer — dérive → POM ; bugs produit → Jira.",
        "Les humains approuvent plan et merge ; les agents apportent la vitesse dans ces rails.",
    ]
    for i, t in enumerate(takeaways):
        top = Inches(1.5) + i * Inches(0.7)
        num = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), top, Inches(0.45), Inches(0.45))
        num.fill.solid()
        num.fill.fore_color.rgb = CGI_RED
        num.line.fill.background()
        tf = num.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tf.paragraphs[0].add_run()
        run.text = str(i + 1)
        set_run(run, size=12, bold=True, color=WHITE)
        tf.paragraphs[0].space_before = Pt(6)
        add_text_box(s, Inches(1.4), top + Inches(0.05), Inches(11), Inches(0.45),
                     t, size=13, color=DARK)
    add_text_box(s, Inches(0.9), Inches(5.3), Inches(11), Inches(0.4),
                 "Questions ?", size=22, bold=True, color=CGI_RED)
    add_text_box(s, Inches(0.9), Inches(5.85), Inches(11), Inches(0.4),
                 "docs/harness-diagram.md · .cursor/rules/ · .cursor/skills/",
                 size=13, color=MUTED)

    out = "/Users/svetlana.sevcenco/Formation/homework/docs/QA-Harness-Techniques-FR.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
